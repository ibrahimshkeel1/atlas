"""Accountant-oriented Excel / PDF export (working papers, not complex GL reports)."""

from __future__ import annotations

import io
from datetime import date, datetime, timezone
from typing import Any, Optional
from uuid import UUID, uuid4

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy.orm import Session

from app.models.organization import Organization
from app.models.report import Report, ReportStatus, ReportType
from app.models.user import User
from app.services.dashboard import report_payload
from app.services.storage import presigned_url, upload_bytes

PKR_FORMAT = '"PKR "#,##0.00'
PCT_FORMAT = "0%"
HEADER_FILL = PatternFill("solid", fgColor="0F172A")
HEADER_FONT = Font(color="FFFFFF", bold=True, name="Calibri", size=11)
LABEL_FONT = Font(bold=True, name="Calibri", size=11)
TITLE_FONT = Font(bold=True, name="Calibri", size=14)
BODY_FONT = Font(name="Calibri", size=11)
THIN = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1"),
)
MUTED_FILL = PatternFill("solid", fgColor="F8FAFC")


def _autosize(ws, min_width: float = 10, max_width: float = 48) -> None:
    for col_cells in ws.columns:
        letter = get_column_letter(col_cells[0].column)
        longest = 0
        for cell in col_cells:
            if cell.value is None:
                continue
            longest = max(longest, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(max(longest + 2, min_width), max_width)


def _write_working_paper_header(
    ws,
    *,
    client_name: str,
    period_start: str,
    period_end: str,
    prepared_date: str,
    prepared_by: str,
    start_row: int = 1,
) -> int:
    """Write standard accountant header. Returns next empty row index."""
    ws.cell(start_row, 1, "Atlas Finance AI — Working papers").font = TITLE_FONT
    rows = [
        ("Client name", client_name),
        ("Period", f"{period_start} to {period_end}"),
        ("Prepared date", prepared_date),
        ("Prepared by", prepared_by),
        ("Currency", "PKR"),
    ]
    r = start_row + 1
    for label, value in rows:
        ws.cell(r, 1, label).font = LABEL_FONT
        ws.cell(r, 2, value).font = BODY_FONT
        r += 1
    return r + 1  # blank row after header


def _apply_table_header(ws, row: int, headers: list[str]) -> None:
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row, col, title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.border = THIN


def build_excel(
    payload: dict,
    *,
    client_name: str,
    prepared_by: str,
    prepared_date: Optional[str] = None,
) -> bytes:
    prepared_date = prepared_date or datetime.now(timezone.utc).strftime("%Y-%m-%d")
    period_start = payload["period_start"]
    period_end = payload["period_end"]

    wb = Workbook()

    # ----- Transactions sheet -----
    ws = wb.active
    ws.title = "Transactions"
    next_row = _write_working_paper_header(
        ws,
        client_name=client_name,
        period_start=period_start,
        period_end=period_end,
        prepared_date=prepared_date,
        prepared_by=prepared_by,
    )

    headers = [
        "Date",
        "Description",
        "Debit",
        "Credit",
        "Category",
        "Confidence",
        "Review status",
    ]
    header_row = next_row
    _apply_table_header(ws, header_row, headers)

    data_start = header_row + 1
    row = data_start
    for t in payload["transactions"]:
        ws.cell(row, 1, t["date"]).font = BODY_FONT
        ws.cell(row, 2, t["description"]).font = BODY_FONT

        debit_cell = ws.cell(row, 3, t["debit"] if t["debit"] is not None else None)
        credit_cell = ws.cell(row, 4, t["credit"] if t["credit"] is not None else None)
        debit_cell.number_format = PKR_FORMAT
        credit_cell.number_format = PKR_FORMAT
        debit_cell.font = BODY_FONT
        credit_cell.font = BODY_FONT

        ws.cell(row, 5, t.get("category") or "Uncategorized").font = BODY_FONT

        conf = t.get("confidence")
        conf_cell = ws.cell(row, 6, conf if conf is not None else None)
        if conf is not None:
            conf_cell.number_format = PCT_FORMAT
        conf_cell.font = BODY_FONT

        ws.cell(row, 7, t.get("review_status") or "—").font = BODY_FONT

        for col in range(1, 8):
            ws.cell(row, col).border = THIN
            if row % 2 == 0:
                ws.cell(row, col).fill = MUTED_FILL
        row += 1

    data_end = row - 1
    # Totals row
    if data_end >= data_start:
        total_row = row
        ws.cell(total_row, 1, "").border = THIN
        ws.cell(total_row, 2, "Totals").font = LABEL_FONT
        ws.cell(total_row, 2).border = THIN
        debit_total = ws.cell(
            total_row, 3, f"=SUM(C{data_start}:C{data_end})"
        )
        credit_total = ws.cell(
            total_row, 4, f"=SUM(D{data_start}:D{data_end})"
        )
        debit_total.number_format = PKR_FORMAT
        credit_total.number_format = PKR_FORMAT
        debit_total.font = LABEL_FONT
        credit_total.font = LABEL_FONT
        for col in range(3, 8):
            ws.cell(total_row, col).border = THIN
            if col >= 5:
                ws.cell(total_row, col).font = BODY_FONT

    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = f"A{header_row}:G{max(data_end, header_row)}"
    _autosize(ws)
    ws.column_dimensions["B"].width = 42

    # ----- Summary sheet -----
    summary = wb.create_sheet("Summary", 1)
    next_row = _write_working_paper_header(
        summary,
        client_name=client_name,
        period_start=period_start,
        period_end=period_end,
        prepared_date=prepared_date,
        prepared_by=prepared_by,
    )

    summary.cell(next_row, 1, "Income / expense overview").font = LABEL_FONT
    next_row += 1
    overview_header = next_row
    _apply_table_header(summary, overview_header, ["Metric", "Amount (PKR)"])
    overview = [
        ("Total income (credits)", payload["total_income"]),
        ("Total expenses (debits)", payload["total_expenses"]),
        ("Net cash flow", payload["net_cash_flow"]),
        ("Transaction count", payload["transaction_count"]),
        ("Reviewed", payload.get("reviewed_count", 0)),
        ("Needs review", payload.get("needs_review_count", 0)),
    ]
    r = overview_header + 1
    for label, value in overview:
        summary.cell(r, 1, label).font = BODY_FONT
        summary.cell(r, 1).border = THIN
        cell = summary.cell(r, 2, value)
        cell.border = THIN
        cell.font = BODY_FONT
        if isinstance(value, float):
            cell.number_format = PKR_FORMAT
        r += 1

    r += 1
    summary.cell(r, 1, "Category totals").font = LABEL_FONT
    r += 1

    income_cats = sorted(
        [c for c in payload["by_category"] if c.get("type") == "income"],
        key=lambda c: c["name"],
    )
    expense_cats = sorted(
        [c for c in payload["by_category"] if c.get("type") != "income"],
        key=lambda c: c["name"],
    )

    def write_category_block(title: str, cats: list[dict[str, Any]]) -> None:
        nonlocal r
        summary.cell(r, 1, title).font = LABEL_FONT
        r += 1
        _apply_table_header(
            summary,
            r,
            ["Category", "Type", "Income (PKR)", "Expenses (PKR)", "Net (PKR)"],
        )
        r += 1
        if not cats:
            summary.cell(r, 1, "— None —").font = BODY_FONT
            r += 2
            return
        for c in cats:
            income_v = float(c.get("income") or 0)
            expense_v = float(c.get("expenses") or 0)
            summary.cell(r, 1, c["name"]).font = BODY_FONT
            summary.cell(r, 2, c.get("type") or "expense").font = BODY_FONT
            for col, val in enumerate(
                [income_v, expense_v, income_v - expense_v], start=3
            ):
                cell = summary.cell(r, col, val)
                cell.number_format = PKR_FORMAT
                cell.font = BODY_FONT
            for col in range(1, 6):
                summary.cell(r, col).border = THIN
            r += 1
        r += 1

    write_category_block("Income categories", income_cats)
    write_category_block("Expense categories", expense_cats)

    summary.freeze_panes = "A8"
    _autosize(summary)
    summary.column_dimensions["A"].width = 28

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_pdf(
    payload: dict,
    *,
    client_name: str,
    prepared_by: str,
    prepared_date: Optional[str] = None,
) -> bytes:
    prepared_date = prepared_date or datetime.now(timezone.utc).strftime("%Y-%m-%d")
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter)
    styles = getSampleStyleSheet()
    story = [
        Paragraph("Atlas Finance AI — Working papers", styles["Title"]),
        Spacer(1, 8),
        Paragraph(f"Client name: {client_name}", styles["Normal"]),
        Paragraph(
            f"Period: {payload['period_start']} to {payload['period_end']}",
            styles["Normal"],
        ),
        Paragraph(f"Prepared date: {prepared_date}", styles["Normal"]),
        Paragraph(f"Prepared by: {prepared_by}", styles["Normal"]),
        Paragraph("Currency: PKR", styles["Normal"]),
        Spacer(1, 12),
        Paragraph(
            f"Total income: PKR {payload['total_income']:,.2f}",
            styles["Normal"],
        ),
        Paragraph(
            f"Total expenses: PKR {payload['total_expenses']:,.2f}",
            styles["Normal"],
        ),
        Paragraph(
            f"Net cash flow: PKR {payload['net_cash_flow']:,.2f}",
            styles["Normal"],
        ),
        Spacer(1, 18),
        Paragraph("Category totals", styles["Heading2"]),
    ]

    table_data = [["Category", "Type", "Income (PKR)", "Expenses (PKR)"]]
    for c in sorted(payload["by_category"], key=lambda x: x["name"]):
        table_data.append(
            [
                c["name"],
                c["type"],
                f"{c['income']:,.2f}",
                f"{c['expenses']:,.2f}",
            ]
        )
    table = Table(table_data, hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F172A")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("BOTTOMPADDING", (0, 0), (-1, 0), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 1), (-1, -1), 4),
            ]
        )
    )
    story.append(table)
    doc.build(story)
    return buf.getvalue()


def create_export(
    db: Session,
    *,
    organization_id: UUID,
    report_type: str,
    period_start: date,
    period_end: date,
    fmt: str,
    prepared_by_user_id: Optional[UUID] = None,
) -> dict:
    payload = report_payload(db, organization_id, period_start, period_end)

    org = db.get(Organization, organization_id)
    client_name = org.name if org else "Client"

    prepared_by = "Atlas"
    if prepared_by_user_id:
        user = db.get(User, prepared_by_user_id)
        if user:
            prepared_by = user.full_name or user.email

    prepared_date = datetime.now(timezone.utc).strftime("%Y-%m-%d")

    try:
        rtype = ReportType(report_type)
    except ValueError:
        rtype = ReportType.CASH_FLOW

    if fmt == "pdf":
        data = build_pdf(
            payload,
            client_name=client_name,
            prepared_by=prepared_by,
            prepared_date=prepared_date,
        )
        content_type = "application/pdf"
        ext = "pdf"
    else:
        data = build_excel(
            payload,
            client_name=client_name,
            prepared_by=prepared_by,
            prepared_date=prepared_date,
        )
        content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ext = "xlsx"

    key = f"orgs/{organization_id}/reports/{uuid4()}.{ext}"
    upload_bytes(data, key, content_type)

    report = Report(
        organization_id=organization_id,
        type=rtype,
        period_start=period_start,
        period_end=period_end,
        s3_key=key,
        status=ReportStatus.READY,
        params_json={
            "format": fmt,
            "client_name": client_name,
            "prepared_by": prepared_by,
            "prepared_date": prepared_date,
        },
        format=fmt,
    )
    db.add(report)
    db.commit()
    db.refresh(report)

    return {
        "id": str(report.id),
        "download_url": presigned_url(key),
        "summary": {
            "total_income": payload["total_income"],
            "total_expenses": payload["total_expenses"],
            "net_cash_flow": payload["net_cash_flow"],
            "transaction_count": payload["transaction_count"],
            "by_category": payload["by_category"],
            "tax_related_count": len(payload["tax_related"]),
            "client_name": client_name,
            "prepared_by": prepared_by,
            "prepared_date": prepared_date,
        },
    }
